from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from app.chunker import SourceDocument, chunk_document
from app.ingest import _requested_file_types, _validate_documents
from app.parser_excel import parse_excel_rows
from app.parser_freshdesk import _build_ticket_metadata
from app.parser_json import parse_stackoverflow_json_files
from app.parser_md import parse_markdown_files


def test_markdown_ingest_unchanged(tmp_path: Path) -> None:
    md_file = tmp_path / "guide.md"
    md_file.write_text("# Hello\n\n## Section\n\nThis is markdown content.", encoding="utf-8")

    docs = parse_markdown_files(tmp_path, only_file_path=str(md_file))
    assert len(docs) == 1
    assert docs[0].source_type == "md"
    assert "markdown content" in docs[0].content


def test_stackoverflow_json_path_still_works(tmp_path: Path) -> None:
    posts = [
        {
            "id": 10,
            "postTypeId": 1,
            "title": "Question title",
            "body": "<p>Question body</p>",
            "creationDate": "2026-01-01T00:00:00Z",
        },
        {
            "id": 11,
            "postTypeId": 2,
            "parentId": 10,
            "body": "<p>Answer body</p>",
        },
    ]
    posts_path = tmp_path / "posts.json"
    posts_path.write_text(__import__("json").dumps(posts), encoding="utf-8")

    docs = parse_stackoverflow_json_files(tmp_path, only_file_path=str(posts_path))
    assert len(docs) == 1
    assert docs[0].source_id == "10"
    assert "Question Title" in docs[0].content


def test_generic_json_export_yields_multiple_docs(tmp_path: Path) -> None:
    payload = {
        "tickets": [
            {"ticket_id": 1, "subject": "Video issue A", "status": "closed"},
            {"ticket_id": 2, "subject": "Video issue B", "status": "open"},
        ]
    }
    json_path = tmp_path / "export.json"
    json_path.write_text(__import__("json").dumps(payload), encoding="utf-8")

    docs = parse_stackoverflow_json_files(tmp_path, only_file_path=str(json_path))
    assert len(docs) == 2
    assert {doc.source_id for doc in docs} == {"1", "2"}


def test_xlsx_parsing_unchanged(tmp_path: Path) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Sheet1"
    ws.append(["ID", "Issue"])
    ws.append([101, "Incomplete video"])
    xlsx_path = tmp_path / "cases.xlsx"
    workbook.save(xlsx_path)

    docs, summary = parse_excel_rows(tmp_path, str(xlsx_path))
    assert len(docs) == 1
    assert docs[0].source_type == "excel"
    assert summary["detected_file_type"] == ".xlsx"


def test_csv_parsing_works(tmp_path: Path) -> None:
    csv_path = tmp_path / "cases.csv"
    csv_path.write_text("ID,Issue\n201,Blackout\n", encoding="utf-8")

    docs, summary = parse_excel_rows(tmp_path, str(csv_path))
    assert len(docs) == 1
    assert docs[0].source_type == "excel"
    assert summary["detected_file_type"] == ".csv"
    assert summary["sheet"] == "csv"


def test_xlsx_all_sheets_mode(tmp_path: Path) -> None:
    workbook = Workbook()
    ws1 = workbook.active
    ws1.title = "One"
    ws1.append(["ID", "Issue"])
    ws1.append([1, "A"])
    ws2 = workbook.create_sheet("Two")
    ws2.append(["ID", "Issue"])
    ws2.append([2, "B"])
    xlsx_path = tmp_path / "multi.xlsx"
    workbook.save(xlsx_path)

    docs, summary = parse_excel_rows(tmp_path, str(xlsx_path), ingest_all_sheets=True)
    assert len(docs) == 2
    assert summary["sheet_count"] == 2


def test_xls_graceful_warning_when_dependency_missing(tmp_path: Path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    import app.parser_excel as parser_excel

    xls_path = tmp_path / "legacy.xls"
    xls_path.write_text("placeholder", encoding="utf-8")
    monkeypatch.setattr(parser_excel, "pd", None)

    docs, summary = parse_excel_rows(tmp_path, str(xls_path))
    assert docs == []
    assert any("pandas is not installed" in warning for warning in summary["warnings"])


def test_file_type_routing_for_tabular_extensions() -> None:
    assert _requested_file_types("cases.xlsx") == (False, False, True)
    assert _requested_file_types("cases.xls") == (False, False, True)
    assert _requested_file_types("cases.csv") == (False, False, True)


def test_deterministic_dedupe_uses_stable_hash() -> None:
    documents = [
        SourceDocument(
            source_type="json",
            source_id="abc",
            content="same content",
            title="A",
            heading=None,
            tags=[],
            creation_date=None,
            metadata={},
        ),
        SourceDocument(
            source_type="json",
            source_id="abc",
            content="same content",
            title="A",
            heading=None,
            tags=[],
            creation_date=None,
            metadata={},
        ),
    ]
    accepted, _, duplicates = _validate_documents(
        documents=documents,
        ingest_run_id="run",
        parser_version="v1",
        tenant=None,
        access_scope=None,
        min_document_chars=1,
    )
    assert len(accepted) == 1
    assert duplicates == 1


def test_freshdesk_metadata_is_compact() -> None:
    metadata = _build_ticket_metadata(
        ticket={"requester_id": 1, "responder_id": 2, "group_id": 3},
        ticket_id_str="123",
        ticket_url="https://example.freshdesk.com/a/tickets/123",
        status="closed",
        priority="high",
        tags=["video"],
        created_at="2026-01-01T00:00:00Z",
        updated_at="2026-01-01T01:00:00Z",
        ticket_type="incident",
        ingest_run_ts="2026-01-01T02:00:00Z",
        freshdesk_sync_cursor=None,
        image_urls=["https://example.com/a.png"],
        conversations=[{"id": 1, "private": False}],
    )
    assert "ticket_fields" not in metadata
    assert "conversations" not in metadata
    assert metadata["raw_payload_ref"] == "freshdesk-ticket:123"


def test_chunker_does_not_drop_large_single_section_content() -> None:
    long_text = " ".join(f"word{i}" for i in range(1300))
    doc = SourceDocument(
        source_type="md",
        source_id="large-doc",
        content=f"# Title\n\n## BigSection\n\n{long_text}",
        title="Large Doc",
        heading="BigSection",
        tags=[],
        creation_date=None,
        metadata={"file_path": "large.md"},
    )
    chunks = chunk_document(
        doc,
        min_words=50,
        target_words=200,
        max_words=250,
        overlap_words=0,
        strategy="heading_aware",
        max_chars=5000,
        min_orphan_words=20,
    )
    assert len(chunks) >= 5
    total_words = sum(chunk.word_count for chunk in chunks)
    # All words from the large section should be preserved across chunks.
    assert total_words >= 1300

